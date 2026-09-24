#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成 examples/sample_input.xlsx（虚构演示数据，可随仓库公开）。

覆盖 amount_unifier 各类输入形态，便于本地跑通与测试：
  纯数值、文本 元/万元/万/亿元/千元、公式、区间、长文本、纵向合并单元格。

运行：python examples/make_sample.py
"""
import os

import openpyxl
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "sample_input.xlsx")


def build():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1 (2)"

    headers = ["序号", "企业", "资产名称", "面积(㎡)", "类型", "保险公司", "期限",
               "金额（万元）", "费用（万元）", "备注"]

    # 第1行：标题（合并，模拟真实表）
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1, "示例数据（虚构，仅用于演示）").font = Font(bold=True, size=12)

    # 第2行：表头
    for c, h in enumerate(headers, start=1):
        ws.cell(2, c, h).font = Font(bold=True)

    # 数据行：H=金额(列8)，I=费用(列9)
    rows = [
        (1, "示例A", "办公楼", 1000, "企财", "甲公司", "1年", 11.44, 0.0057, "纯数值"),
        (2, "示例A", "设备", 200, "企财", "甲公司", "1年", 1567.9341, 0.7839, "纯数值"),
        (3, "示例B", "仓库", 500, "企财", "乙公司", "1年", "2004279.6元", "2204.71元", "文本元"),
        (4, "示例B", "商场", 800, "公众责任", "乙公司", "1年", "6581万元", "0.40万元", "文本万元"),
        (5, "示例C", "场馆", 300, "企财", "丙公司", "1年", "777.3万", "0.7773万", "文本万"),
        (6, "示例C", "园区", 1200, "企财", "丙公司", "1年", "2亿元", None, "文本亿元"),
        (7, "示例D", "车库", 400, "责任", "丁公司", "1年", "5千元", None, "文本千元"),
        (8, "示例D", "机房", 100, "企财", "丁公司", "1年", "=11000000+2400000", None, "公式(无缓存→需人工)"),
        (9, "示例E", "多处在建", 0, "工程", "戊公司", "1年", "1万-100万不等", None, "区间→需人工"),
        (10, "示例E", "综合店", 600, "多险种", "戊公司", "1年",
            "累计赔偿限额200万，每次事故100万，附加责任若干……", None, "长文本→需人工"),
    ]
    r = 3
    for row in rows:
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(r, c, v)
        r += 1

    # 合并示例：H13:H15 共用一额度（锚点 H13，H14/H15 继承）
    ws.cell(13, 1, 11); ws.cell(13, 3, "合并资产甲"); ws.cell(13, 8, "973.7256万元")
    ws.cell(14, 3, "合并资产乙")
    ws.cell(15, 3, "合并资产丙")
    ws.merge_cells(start_row=13, start_column=8, end_row=15, end_column=8)

    wb.save(OUT)
    print("已生成:", OUT)


if __name__ == "__main__":
    build()
